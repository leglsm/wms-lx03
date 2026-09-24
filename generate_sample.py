"""Generate fictional SAP LX03-style bin stock sample data for the WMS LX03 Viewer demo.

All names, plants, materials and quantities are fictional. Fixed seed -> same output every run.
Usage: python scripts/generate_sample.py   (writes data/sample_lx03.xlsx and data/sample.json)
"""
import json
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

SEED = 20260924
AS_OF = date(2026, 9, 1)          # fixed anchor so aging is reproducible
WAREHOUSE = "W10"
PLANT = "US10"                    # fictional plant code

# Storage type layout: (type, aisles, bays, levels, fill rate)
LAYOUT = [
    ("001", 6, 20, 4, 0.78),      # standard pallet racking
    ("002", 3, 15, 3, 0.70),      # overflow racking
    ("003", 2, 10, 5, 0.65),      # highbay
]
FLOOR_BINS = [f"FLR-{i:02d}" for i in range(1, 13)]      # 900 floor staging (free-form)

# Fictional material families (SAP-style number ranges)
FAMILIES = [
    # prefix, start, count, UoM, description stems, qty range per HU
    ("2", 20410000, 55, "EA", ["BRACKET", "CLIP", "RETAINER", "GRILLE INSERT", "SENSOR HOLDER",
                               "REINFORCEMENT", "SEAL STRIP", "FOG LAMP BEZEL"], (40, 400)),
    ("5", 51820000, 30, "EA", ["FASCIA MOLDED", "SPOILER MOLDED", "ABSORBER", "SKID PLATE"], (8, 48)),
    ("45", 45610000, 25, "EA", ["FRONT FASCIA ASSY", "REAR FASCIA ASSY", "SIDE SILL ASSY",
                                "LIFTGATE TRIM ASSY"], (4, 24)),
    ("6", 61200000, 10, "EA", ["RETURNABLE RACK", "PLASTIC TOTE 24x15", "DUNNAGE TRAY"], (1, 20)),
]
PROGRAMS = ["NX1", "NX2", "KT4"]       # fictional vehicle programs
SIDES = ["LH", "RH", ""]


def build_materials(rng):
    mats = []
    for _, start, count, uom, stems, qty_rng in FAMILIES:
        for i in range(count):
            num = str(start + i * rng.randint(3, 17))
            desc = f"{rng.choice(stems)} {rng.choice(PROGRAMS)} {rng.choice(SIDES)}".strip()
            mats.append({"mat": num, "desc": desc, "uom": uom, "qty": qty_rng,
                         "slow": rng.random() < 0.12})   # ~12% slow movers -> old GR dates
    return mats


def gr_date(rng, slow):
    days = rng.randint(200, 540) if slow else int(rng.expovariate(1 / 25))
    return AS_OF - timedelta(days=min(days, 540))


def main():
    rng = random.Random(SEED)
    mats = build_materials(rng)
    rows, hu = [], 1000450000

    bins = []
    for stype, aisles, bays, levels, fill in LAYOUT:
        for a in range(1, aisles + 1):
            for b in range(1, bays + 1):
                for lv in range(1, levels + 1):
                    if rng.random() < fill:
                        bins.append((stype, f"{a:02d}-{b:02d}-{lv:02d}"))
    bins += [("900", fb) for fb in FLOOR_BINS]

    for stype, bin_ in bins:
        n_hu = 1 if stype != "900" else rng.randint(2, 6)
        for _ in range(n_hu):
            m = rng.choice(mats)
            hu += rng.randint(1, 9)
            qty = rng.randint(*m["qty"])
            blocked = rng.random() < 0.04
            rows.append({
                "Warehouse No.": WAREHOUSE,
                "Storage Type": stype,
                "Storage Bin": bin_,
                "Storage Unit": str(hu).zfill(10),
                "Material": m["mat"],
                "Material Description": m["desc"],
                "Plant": PLANT,
                "Batch": f"B{rng.randint(2601, 2635)}",
                "Total Stock": qty,
                "Available Stock": 0 if blocked else qty,
                "Base Unit of Measure": m["uom"],
                "GR Date": gr_date(rng, m["slow"]),
            })

    out = Path(__file__).resolve().parent.parent / "data"
    out.mkdir(exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "LX03"
    headers = list(rows[0].keys())
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r[h] for h in headers])
    for cell in ws["L"][1:]:
        cell.number_format = "mm/dd/yyyy"
    for col, w in zip("ABCDEFGHIJKL", [13, 12, 12, 13, 11, 30, 7, 8, 11, 14, 10, 11]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(out / "sample_lx03.xlsx")

    with open(out / "sample.json", "w") as f:
        json.dump([{**r, "GR Date": r["GR Date"].isoformat()} for r in rows], f, indent=1)

    print(f"{len(rows)} rows, {len(bins)} bins, {len(mats)} materials -> {out}")


if __name__ == "__main__":
    main()
